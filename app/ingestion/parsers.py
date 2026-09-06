import csv
from io import StringIO
from pathlib import Path


class DocumentParser:
    def parse(self, path: Path, content_type: str) -> tuple[str, dict]:
        suffix = path.suffix.lower()
        if suffix == ".pdf" or content_type == "application/pdf":
            return self._parse_pdf(path)
        if suffix == ".docx":
            return self._parse_docx(path)
        if suffix == ".csv":
            return self._parse_csv(path)
        if suffix in {".txt", ".md", ".log"}:
            return path.read_text(encoding="utf-8", errors="ignore"), {"parser": "text"}
        if suffix in {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}:
            return self._parse_image(path)
        if suffix == ".pptx":
            return self._parse_pptx(path)
        return path.read_text(encoding="utf-8", errors="ignore"), {"parser": "fallback_text"}

    def _parse_pdf(self, path: Path) -> tuple[str, dict]:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages), {"parser": "pypdf", "page_count": len(reader.pages)}

    def _parse_docx(self, path: Path) -> tuple[str, dict]:
        from docx import Document

        document = Document(str(path))
        return "\n".join(p.text for p in document.paragraphs), {"parser": "python-docx"}

    def _parse_csv(self, path: Path) -> tuple[str, dict]:
        raw = path.read_text(encoding="utf-8", errors="ignore")
        reader = csv.reader(StringIO(raw))
        lines = [" | ".join(row) for row in reader]
        return "\n".join(lines), {"parser": "csv", "row_count": len(lines)}

    def _parse_image(self, path: Path) -> tuple[str, dict]:
        try:
            from PIL import Image
            import pytesseract

            text = pytesseract.image_to_string(Image.open(path))
            return text, {"parser": "pytesseract", "ocr": True}
        except Exception:
            return "", {"parser": "ocr_unavailable", "ocr": False}

    def _parse_pptx(self, path: Path) -> tuple[str, dict]:
        try:
            from pptx import Presentation

            deck = Presentation(str(path))
            text: list[str] = []
            for index, slide in enumerate(deck.slides, start=1):
                text.append(f"Slide {index}")
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text.append(shape.text)
            return "\n".join(text), {"parser": "python-pptx", "slide_count": len(deck.slides)}
        except Exception:
            return "", {"parser": "pptx_unavailable"}

